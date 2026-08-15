from pathlib import Path

from docx import Document

from dataprocess.file_converter import convert_local_document


def testConvertsUtf8TextFile(tmp_path: Path) -> None:
    source = tmp_path / "guide.md"
    source.write_text("# 标题\n\n内容", encoding="utf-8")

    assert convert_local_document(source) == "# 标题\n\n内容"


def testConvertsDocxParagraphsAndTables(tmp_path: Path) -> None:
    source = tmp_path / "products.docx"
    document = Document()
    document.add_paragraph("商品说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "价格"
    table.cell(1, 0).text = "耳机"
    table.cell(1, 1).text = "99"
    document.save(source)

    converted = convert_local_document(source)

    assert "商品说明" in converted
    assert "| 名称 | 价格 |" in converted
    assert "| 耳机 | 99 |" in converted
