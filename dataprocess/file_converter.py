"""按文件后缀将本地文档转换为 Markdown 文本。"""

from __future__ import annotations

from pathlib import Path
from typing import Protocol

from docx import Document

LOCAL_SUFFIXES = {".txt", ".md", ".docx"}
MINERU_SUFFIXES = {".pdf", ".ppt", ".pptx"}
SUPPORTED_SUFFIXES = LOCAL_SUFFIXES | MINERU_SUFFIXES


class MinerUParser(Protocol):
    """MinerU 适配器的最小接口。"""

    def parse_files(self, file_paths: list[Path]) -> list[str]:
        """解析上传文件，按输入顺序返回 Markdown。"""


def convert_local_document(path: Path) -> str:
    """转换 TXT、MD 或 DOCX；PDF/PPT 由调用方交给 MinerU。"""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _read_text(path)
    if suffix == ".docx":
        return _read_docx(path)
    raise ValueError(f"Unsupported local document suffix: {suffix}")


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("document", b"", 0, 1, f"Unable to decode {path}")


def _read_docx(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = [[cell.text.replace("|", "\\|").strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        blocks.append("| " + " | ".join(rows[0]) + " |")
        blocks.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n\n".join(blocks)
